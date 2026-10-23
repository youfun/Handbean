"""Regenerate synthetic fixtures with production document writers, never at runtime.

Test-only requirements: reportlab==5.1.0 python-docx==1.2.0 openpyxl==3.1.5
python-pptx==1.0.2. Pass the path of LongCang-Regular.ttf as argv[1].
Font source: https://github.com/google/fonts/tree/main/ofl/longcang (OFL.txt).
All document content is synthetic. The PDF embeds a font subset and ToUnicode.
"""
from pathlib import Path
import sys
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pdfencrypt import StandardEncryption
from reportlab.lib.utils import ImageReader
from PIL import Image, ImageDraw
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

root = Path(__file__).parent
pdfmetrics.registerFont(TTFont("FixtureChinese", sys.argv[1]))
pdf = canvas.Canvas(str(root / "multilingual.pdf"), invariant=1, pageCompression=1)
pdf.setFont("Helvetica", 12)
pdf.drawString(72, 750, "Quarterly report: revenue 123.45 EUR")
pdf.setFont("FixtureChinese", 14)
pdf.drawString(72, 720, "中文文档测试：你好世界")
pdf.showPage()
pdf.setFont("Helvetica", 12)
pdf.drawString(72, 750, "Second page: cafe and report conclusion")
pdf.showPage()
pdf.save()
blank = canvas.Canvas(str(root / "image-only.pdf"), invariant=1)
image = Image.new("RGB", (300, 80), "white")
ImageDraw.Draw(image).text((10, 20), "Scanned content requires OCR", fill="black")
blank.drawImage(ImageReader(image), 72, 650, width=300, height=80)
blank.showPage()
blank.save()
encrypted = canvas.Canvas(str(root / "encrypted.pdf"), invariant=1, encrypt=StandardEncryption("test-only-password"))
encrypted.drawString(72, 750, "This file is password-protected")
encrypted.showPage()
encrypted.save()

doc = Document()
doc.add_paragraph("Quarterly report 中文段落")
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Item"
table.cell(0, 1).text = "Amount"
table.cell(1, 0).text = "研发"
table.cell(1, 1).text = "123.45"
doc.add_paragraph("Final paragraph")
doc.save(root / "report.docx")

book = Workbook()
sheet = book.active
sheet.title = "预算"
sheet["A1"] = "Category"
sheet["C1"] = "研发"
sheet["A3"] = 123.45
sheet["C3"] = "=A3*2"
book.create_sheet("Summary")["B2"] = "Final total"
book.save(root / "budget.xlsx")

slides = Presentation()
for title, body in [("Opening 中文", "First slide"), ("Conclusion", "Second slide")]:
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
slides.save(root / "slides.pptx")
